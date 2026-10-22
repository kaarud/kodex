def test_build_locus_centers_window():
    from igv_link import build_locus
    assert build_locus("9", 1000, half_window=145) == "chr9:855-1145"

def test_build_locus_already_prefixed():
    from igv_link import build_locus
    assert build_locus("chr16", 2098100, half_window=10) == "chr16:2098090-2098110"

def test_build_locus_clamps_start_to_one():
    from igv_link import build_locus
    assert build_locus("chrX", 50, half_window=145) == "chrX:1-195"


def _touch_run(tmp_path):
    """Create a fake bam_dir with two samples (S1, S2), each with 3 BAMs."""
    for sid in ("100008_S1", "999111_S2"):
        for name in (f"{sid}.bam", f"{sid}.bam.bai",
                     f"{sid}.chim.bam", f"{sid}.chim.bam.bai",
                     f"{sid}.mutect2.bam", f"{sid}.mutect2.bam.bai"):
            (tmp_path / name).touch()
    return tmp_path

def test_resolve_reference_bam_by_lane(tmp_path):
    from igv_link import resolve_reference_bam
    bam_dir = _touch_run(tmp_path)
    ref = resolve_reference_bam(bam_dir, "S2", current_sample_id="100008_S1")
    assert ref.name == "999111_S2.bam"

def test_resolve_reference_bam_falls_back_when_current_is_reference(tmp_path):
    """Comparing the reference sample to itself is meaningless → another sample."""
    from igv_link import resolve_reference_bam
    bam_dir = _touch_run(tmp_path)
    ref = resolve_reference_bam(bam_dir, "S2", current_sample_id="999111_S2")
    assert ref.name == "100008_S1.bam"

def test_resolve_reference_bam_missing_lane_falls_back(tmp_path):
    from igv_link import resolve_reference_bam
    bam_dir = _touch_run(tmp_path)
    ref = resolve_reference_bam(bam_dir, "S9", current_sample_id="100008_S1")
    assert ref.name == "999111_S2.bam"

def test_resolve_reference_bam_none_when_alone(tmp_path):
    """A single-sample run has nothing to compare against."""
    from igv_link import resolve_reference_bam
    for name in ("100008_S1.bam", "100008_S1.chim.bam", "100008_S1.mutect2.bam"):
        (tmp_path / name).touch()
    assert resolve_reference_bam(tmp_path, "S2", current_sample_id="100008_S1") is None

def test_resolve_reference_bam_ignores_chim_and_mutect2(tmp_path):
    """The raw .bam is returned, never the .chim.bam / .mutect2.bam of the lane."""
    from igv_link import resolve_reference_bam
    bam_dir = _touch_run(tmp_path)
    ref = resolve_reference_bam(bam_dir, "S2", current_sample_id="100008_S1")
    assert ref.name == "999111_S2.bam"
    assert ".chim." not in ref.name and ".mutect2." not in ref.name


def test_collect_bam_tracks_patient_plus_reference(tmp_path):
    from igv_link import collect_bam_tracks
    bam_dir = _touch_run(tmp_path)   # S1 + S2 present
    tracks = collect_bam_tracks("100008_S1", bam_dir, reference_lane="S2")
    labels = [label for _, label in tracks]
    names = [p.name for p, _ in tracks]
    assert labels == ["RAW", "CHIM", "MUTECT2", "REF_S2"]
    assert names == ["100008_S1.bam", "100008_S1.chim.bam",
                     "100008_S1.mutect2.bam", "999111_S2.bam"]

def test_collect_bam_tracks_fallback_reference_when_self(tmp_path):
    """The lane sample is the patient itself → another sample, labelled by name."""
    from igv_link import collect_bam_tracks
    bam_dir = _touch_run(tmp_path)
    tracks = collect_bam_tracks("999111_S2", bam_dir, reference_lane="S2")
    labels = [label for _, label in tracks]
    assert labels == ["RAW", "CHIM", "MUTECT2", "REF_100008_S1"]
    assert tracks[-1][0].name == "100008_S1.bam"

def test_collect_bam_tracks_returns_none_when_patient_bams_missing(tmp_path):
    from igv_link import collect_bam_tracks
    # no raw .bam for the patient → None
    assert collect_bam_tracks("100008_S1", tmp_path, reference_lane="S2") is None


def test_collect_bam_tracks_optional_types_raw_and_chim_only(tmp_path):
    """Run 2 ships raw + chim only (no mutect2): load what exists."""
    from igv_link import collect_bam_tracks
    # patient 100005_S1 has raw + chim; reference 100004_S2 has raw + chim
    for sid in ("100005_S1", "100004_S2"):
        for name in (f"{sid}.bam", f"{sid}.bai",
                     f"{sid}.chim.bam", f"{sid}.chim.bai"):
            (tmp_path / name).touch()
    tracks = collect_bam_tracks("100005_S1", tmp_path, reference_lane="S2")
    labels = [label for _, label in tracks]
    names = [p.name for p, _ in tracks]
    assert labels == ["RAW", "CHIM", "REF_S2"]   # no MUTECT2 track
    assert names == ["100005_S1.bam", "100005_S1.chim.bam", "100004_S2.bam"]


def test_collect_bam_tracks_raw_only(tmp_path):
    """Raw alone is enough: patient track loads, no chim/mutect2."""
    from igv_link import collect_bam_tracks
    (tmp_path / "100005_S1.bam").touch()
    tracks = collect_bam_tracks("100005_S1", tmp_path, reference_lane="S2")
    assert [label for _, label in tracks] == ["RAW"]


from pathlib import Path as _P

def test_build_load_url_full():
    from igv_link import build_load_url
    tracks = [
        (_P("/data/100008_S1.bam"), "RAW"),
        (_P("/data/100008_S1.chim.bam"), "CHIM"),
        (_P("/data/100008_S1.mutect2.bam"), "MUTECT2"),
        (_P("/data/999111_S2.bam"), "REF_S2"),
    ]
    url = build_load_url(tracks, locus="chr16:2098000-2098290",
                         port=60151, genome="hg38")
    assert url == (
        "http://localhost:60151/load?"
        "file=/data/100008_S1.bam,/data/100008_S1.chim.bam,"
        "/data/100008_S1.mutect2.bam,/data/999111_S2.bam"
        "&genome=hg38&merge=false"
        "&locus=chr16%3A2098000-2098290"
        "&name=RAW,CHIM,MUTECT2,REF_S2"
    )

def test_build_load_url_encodes_spaces_in_path():
    from igv_link import build_load_url
    tracks = [(_P("/da ta/s.bam"), "RAW")]
    url = build_load_url(tracks, locus="chr9:1-2", port=60151, genome="hg38")
    assert "/da%20ta/s.bam" in url
    assert "file=/da%20ta/s.bam&" in url

def test_build_load_url_omits_locus_when_none():
    from igv_link import build_load_url
    url = build_load_url([(_P("/d/a.bam"), "RAW")], locus=None,
                         port=60151, genome="hg38")
    assert "&locus=" not in url
    assert url == ("http://localhost:60151/load?file=/d/a.bam"
                   "&genome=hg38&merge=false&name=RAW")

def test_build_goto_url():
    from igv_link import build_goto_url
    assert (build_goto_url("chr16:2098000-2098290", port=60151)
            == "http://localhost:60151/goto?locus=chr16%3A2098000-2098290")

def test_build_sort_base_url():
    from igv_link import build_sort_base_url
    assert (build_sort_base_url(port=60151)
            == "http://localhost:60151/execute?command=sort%20base")


import openpyxl

def test_inject_igv_column_adds_hyperlinks(tmp_path):
    from igv_link import inject_igv_column
    bam_dir = _touch_run(tmp_path)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["SYMBOL", "CHR", "POS"])
    ws.append(["TSC1", "chr9", 1000])
    ws.append(["TSC2", "chr16", 2098100])

    inject_igv_column(ws, sample_id="100008_S1", bam_dir=bam_dir,
                      reference_lane="S2", port=60151, half_window=145)

    hdr = ws.cell(row=1, column=1)
    assert hdr.value == "⇅ base"           # A1 = sort-by-base button
    assert hdr.hyperlink.target == "http://localhost:60151/execute?command=sort%20base"
    link1 = ws.cell(row=2, column=1).hyperlink.target
    assert link1.startswith("http://localhost:60151/goto?")
    assert "locus=chr9%3A855-1145" in link1
    assert "file=" not in link1            # goto only — no BAM reload per click
    link2 = ws.cell(row=3, column=1).hyperlink.target
    assert "locus=chr16%3A2097955-2098245" in link2

def test_inject_igv_column_noop_when_bams_missing(tmp_path):
    from igv_link import inject_igv_column
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["SYMBOL", "CHR", "POS"])
    ws.append(["TSC1", "chr9", 1000])
    added = inject_igv_column(ws, sample_id="100008_S1", bam_dir=tmp_path,
                              reference_lane="S2", port=60151, half_window=145)
    assert added is False
    assert ws.cell(row=1, column=1).value == "SYMBOL"


def test_inject_summary_load_link_appends_hyperlink(tmp_path):
    from igv_link import inject_summary_load_link
    bam_dir = _touch_run(tmp_path)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Key", "Value"])
    ws.append(["date", "2026-06-17"])

    added = inject_summary_load_link(ws, sample_id="100008_S1", bam_dir=bam_dir,
                                     reference_lane="S2", port=60151, genome="hg38")
    assert added is True
    cell = next(ws.cell(row=r, column=2) for r in range(1, ws.max_row + 1)
                if ws.cell(row=r, column=1).value == "IGV")
    target = cell.hyperlink.target
    assert target.startswith("http://localhost:60151/load?file=")
    assert "name=RAW,CHIM,MUTECT2,REF_S2" in target
    assert "&locus=" not in target          # load only — navigation is per ▶

def test_inject_summary_load_link_noop_when_bams_missing(tmp_path):
    from igv_link import inject_summary_load_link
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Key", "Value"])
    added = inject_summary_load_link(ws, sample_id="100008_S1", bam_dir=tmp_path,
                                     reference_lane="S2")
    assert added is False
    assert ws.max_row == 1                   # sheet untouched


def test_inject_aggregate_igv_columns_load_and_goto(tmp_path):
    from igv_link import inject_aggregate_igv_columns
    bam_dir = _touch_run(tmp_path)           # 100008_S1 + 999111_S2, 3 BAMs each
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["SAMPLE", "SYMBOL", "CHR", "POS"])
    ws.append(["100008_S1", "TSC1", "chr9", 1000])
    ws.append(["999111_S2", "TSC2", "chr16", 2098100])

    n = inject_aggregate_igv_columns(ws, bam_dir=bam_dir, reference_lane="S2",
                                     port=60151, genome="hg38", half_window=145)
    assert n == 2

    # A1 = sort by base ; B1 = goto header
    a1 = ws.cell(row=1, column=1)
    assert a1.value == "⇅ base"
    assert a1.hyperlink.target == "http://localhost:60151/execute?command=sort%20base"
    assert ws.cell(row=1, column=2).value == "goto"

    # Column A (load): per-row sample BAMs + jump to locus
    load = ws.cell(row=2, column=1)
    assert load.value == "▶ BAM"
    assert load.hyperlink.target.startswith("http://localhost:60151/load?file=")
    assert "100008_S1.bam" in load.hyperlink.target
    assert "locus=chr9%3A855-1145" in load.hyperlink.target

    # Column B (goto): navigation only, no BAM reload
    goto = ws.cell(row=2, column=2)
    assert goto.value == "▶"
    assert goto.hyperlink.target.startswith("http://localhost:60151/goto?")
    assert "file=" not in goto.hyperlink.target

    # Second row resolves the other sample
    assert "999111_S2.bam" in ws.cell(row=3, column=1).hyperlink.target

    # Original data shifted right by 2
    assert ws.cell(row=1, column=3).value == "SAMPLE"


def test_inject_aggregate_igv_columns_blank_row_when_sample_bams_missing(tmp_path):
    from igv_link import inject_aggregate_igv_columns
    bam_dir = _touch_run(tmp_path)           # only 100008_S1 / 999111_S2 exist
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["SAMPLE", "CHR", "POS"])
    ws.append(["100008_S1", "chr9", 1000])
    ws.append(["000000_S9", "chr1", 500])    # no BAMs for this sample

    n = inject_aggregate_igv_columns(ws, bam_dir=bam_dir, reference_lane="S2")
    assert n == 1
    assert ws.cell(row=2, column=1).hyperlink is not None    # has BAMs
    assert ws.cell(row=3, column=1).value is None            # blank, no BAMs
    assert ws.cell(row=3, column=1).hyperlink is None


def test_inject_aggregate_igv_columns_noop_without_headers(tmp_path):
    from igv_link import inject_aggregate_igv_columns
    bam_dir = _touch_run(tmp_path)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["SYMBOL", "CHR", "POS"])      # no SAMPLE column
    ws.append(["TSC1", "chr9", 1000])
    n = inject_aggregate_igv_columns(ws, bam_dir=bam_dir, reference_lane="S2")
    assert n == 0
    assert ws.cell(row=1, column=1).value == "SYMBOL"    # untouched


def test_resolve_bam_dir_relative_follows_moved_run(tmp_path):
    """A relative bam_dir resolves against the xlsx dir, so moving the whole
    run directory keeps the BAMs findable."""
    from igv_link import resolve_bam_dir
    run = tmp_path / "kodex_runs" / "run_2"
    (run / "xlsx").mkdir(parents=True)
    (run / "BAM").mkdir()
    assert resolve_bam_dir("../BAM", run / "xlsx") == (run / "BAM").resolve()

def test_resolve_bam_dir_absolute_unchanged(tmp_path):
    from igv_link import resolve_bam_dir
    assert resolve_bam_dir(str(tmp_path), tmp_path / "elsewhere") == tmp_path

def test_resolve_bam_dir_expands_tilde(tmp_path):
    from igv_link import resolve_bam_dir
    from pathlib import Path
    assert resolve_bam_dir("~/data/x", tmp_path) == Path.home() / "data" / "x"
